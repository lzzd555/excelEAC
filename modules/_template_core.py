"""
Excel 模板共享核心原语
公式解析/改写、sheet 匹配、样式/颜色复制、外部链接、模板读取。
新老两个模板方法(template_generator / template_formula)都从此导入。
"""

import os
import re
import copy
from typing import List, Dict, Any, Optional, Tuple
from dataclasses import dataclass

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection


# ==================== 数据类定义 ====================

@dataclass
class CellStyle:
    """单元格样式"""
    font: Optional[Dict[str, Any]] = None
    fill: Optional[Dict[str, Any]] = None
    border: Optional[Dict[str, Any]] = None
    alignment: Optional[Dict[str, Any]] = None
    number_format: Optional[str] = None
    protection: Optional[Dict[str, Any]] = None


# ==================== 颜色处理函数 ====================

def copy_color(source_color):
    """
    复制颜色对象，支持 RGB、主题、索引和自动颜色

    Args:
        source_color: 源颜色对象 (openpyxl.styles.colors.Color)

    Returns:
        新的 Color 对象，或原始颜色值（RGB字符串）
    """
    from openpyxl.styles.colors import Color

    if source_color is None:
        return None

    color_type = getattr(source_color, 'type', None)

    if color_type == 'rgb':
        return _copy_rgb_color(source_color)
    elif color_type == 'theme':
        return _copy_theme_color(source_color)
    elif color_type == 'indexed':
        return _copy_indexed_color(source_color)
    elif color_type == 'auto' or (hasattr(source_color, 'auto') and source_color.auto):
        return Color(auto=True)

    return source_color


def _copy_rgb_color(source_color):
    """复制 RGB 颜色"""
    from openpyxl.styles.colors import Color
    if hasattr(source_color, 'rgb') and isinstance(source_color.rgb, str) and source_color.rgb:
        return source_color.rgb
    return source_color


def _copy_theme_color(source_color):
    """复制主题颜色"""
    from openpyxl.styles.colors import Color
    try:
        theme_val = int(source_color.theme)
        tint_val = float(source_color.tint) if source_color.tint else 0
        return Color(theme=theme_val, tint=tint_val)
    except (TypeError, ValueError):
        return source_color


def _copy_indexed_color(source_color):
    """复制索引颜色"""
    from openpyxl.styles.colors import Color
    try:
        indexed_val = int(source_color.indexed)
        return Color(indexed=indexed_val)
    except (TypeError, ValueError):
        return source_color


def copy_side(source_side):
    """
    复制边框线样式（Side对象），正确处理颜色

    Args:
        source_side: 源边框线对象

    Returns:
        新的 Side 对象
    """
    if source_side is None or source_side.border_style is None:
        return None

    color = copy_color(source_side.color) if source_side.color else None

    return Side(
        style=source_side.border_style,
        color=color
    )


# ==================== 单元格样式复制函数 ====================

def copy_cell_style(source_cell, target_cell) -> None:
    """
    复制单元格样式

    Args:
        source_cell: 源单元格
        target_cell: 目标单元格
    """
    if not source_cell.has_style:
        return

    _copy_font_style(source_cell, target_cell)
    _copy_fill_style(source_cell, target_cell)
    _copy_border_style(source_cell, target_cell)
    _copy_alignment_style(source_cell, target_cell)

    if source_cell.number_format:
        target_cell.number_format = source_cell.number_format

    if source_cell.protection:
        target_cell.protection = Protection(
            locked=source_cell.protection.locked,
            hidden=source_cell.protection.hidden
        )


def _copy_font_style(source_cell, target_cell) -> None:
    """复制字体样式"""
    if not source_cell.font:
        return

    font_args = {
        'name': source_cell.font.name,
        'size': source_cell.font.size,
        'bold': source_cell.font.bold,
        'italic': source_cell.font.italic,
        'vertAlign': source_cell.font.vertAlign,
        'underline': source_cell.font.underline,
        'strike': source_cell.font.strike,
    }

    if source_cell.font.color:
        font_args['color'] = _get_font_color(source_cell.font.color)

    target_cell.font = Font(**font_args)


def _get_font_color(src_color):
    """获取字体颜色"""
    from openpyxl.styles.colors import Color
    color_type = getattr(src_color, 'type', None)

    if color_type == 'rgb' and hasattr(src_color, 'rgb'):
        if isinstance(src_color.rgb, str) and src_color.rgb:
            return Color(rgb=src_color.rgb)
    elif color_type == 'theme' and hasattr(src_color, 'theme'):
        try:
            return Color(theme=int(src_color.theme), tint=float(src_color.tint) if src_color.tint else 0)
        except (TypeError, ValueError):
            pass
    elif color_type == 'indexed' and hasattr(src_color, 'indexed'):
        try:
            return Color(indexed=int(src_color.indexed))
        except (TypeError, ValueError):
            pass
    elif color_type == 'auto' or (hasattr(src_color, 'auto') and src_color.auto):
        return Color(auto=True)

    return src_color


def _copy_fill_style(source_cell, target_cell) -> None:
    """复制填充样式"""
    if not source_cell.fill:
        return

    try:
        fill_type = source_cell.fill.fill_type
        _apply_fill_by_type(source_cell, target_cell, fill_type)
    except Exception as e:
        _apply_fallback_fill(target_cell, fill_type, e)


def _apply_fill_by_type(source_cell, target_cell, fill_type) -> None:
    """根据填充类型应用填充样式"""
    if fill_type is None or fill_type == 'none':
        target_cell.fill = PatternFill(fill_type='none')
    elif fill_type in ('gray125', 'gray0625'):
        target_cell.fill = PatternFill(fill_type=fill_type)
    elif fill_type == 'solid':
        _apply_solid_fill(source_cell, target_cell)
    else:
        target_cell.fill = PatternFill(fill_type=fill_type)


def _apply_solid_fill(source_cell, target_cell) -> None:
    """应用实心填充"""
    from openpyxl.styles.colors import Color
    start_color = source_cell.fill.start_color
    end_color = source_cell.fill.end_color

    if not start_color:
        target_cell.fill = PatternFill(fill_type='none')
        return

    color_value = _get_fill_color_value(start_color, Color)

    if color_value is not None:
        target_cell.fill = PatternFill(
            fill_type='solid',
            start_color=color_value,
            end_color=color_value
        )
    else:
        target_cell.fill = PatternFill(
            fill_type='solid',
            start_color=start_color,
            end_color=end_color if end_color else start_color
        )


def _get_fill_color_value(start_color, Color):
    """获取填充颜色值"""
    if hasattr(start_color, 'rgb') and isinstance(start_color.rgb, str) and start_color.rgb:
        return start_color.rgb
    elif hasattr(start_color, 'theme') and start_color.theme is not None:
        return Color(theme=start_color.theme, tint=start_color.tint or 0)
    elif hasattr(start_color, 'indexed') and start_color.indexed is not None:
        return Color(indexed=start_color.indexed)
    return None


def _apply_fallback_fill(target_cell, fill_type, error) -> None:
    """应用后备填充样式"""
    try:
        if fill_type is None:
            target_cell.fill = PatternFill(fill_type='none')
        else:
            target_cell.fill = PatternFill(fill_type=fill_type)
    except Exception:
        print(f"⚠️ 跳过填充样式复制（错误: {error}）")


def _copy_border_style(source_cell, target_cell) -> None:
    """复制边框样式"""
    if not source_cell.border:
        return

    target_cell.border = Border(
        left=copy_side(source_cell.border.left),
        right=copy_side(source_cell.border.right),
        top=copy_side(source_cell.border.top),
        bottom=copy_side(source_cell.border.bottom),
        diagonal=copy_side(source_cell.border.diagonal),
        diagonal_direction=source_cell.border.diagonal_direction,
        outline=source_cell.border.outline,
        horizontal=copy_side(source_cell.border.horizontal),
        vertical=copy_side(source_cell.border.vertical)
    )


def _copy_alignment_style(source_cell, target_cell) -> None:
    """复制对齐样式"""
    if not source_cell.alignment:
        return

    target_cell.alignment = Alignment(
        horizontal=source_cell.alignment.horizontal,
        vertical=source_cell.alignment.vertical,
        text_rotation=source_cell.alignment.text_rotation,
        wrap_text=source_cell.alignment.wrap_text,
        shrink_to_fit=source_cell.alignment.shrink_to_fit,
        indent=source_cell.alignment.indent
    )


# ==================== 外部链接读取函数 ====================

def read_external_links(xlsx_file: str) -> Dict[int, str]:
    """
    读取Excel文件中的外部链接映射

    Args:
        xlsx_file: Excel文件路径

    Returns:
        Dict[int, str]: 外部链接映射 {索引号: 文件名}
    """
    links = {}

    try:
        with openpyxl.load_workbook(xlsx_file, data_only=False) as wb:
            links = _extract_external_links(wb)
    except Exception as e:
        print(f"   警告: 无法读取外部链接: {e}")

    return links


def _extract_external_links(wb) -> Dict[int, str]:
    """从工作簿中提取外部链接"""
    links = {}

    if not hasattr(wb, 'external_links') or not wb.external_links:
        return links

    for link in wb.external_links:
        link_info = _parse_single_link(link)
        if link_info:
            links.update(link_info)

    return links


def _parse_single_link(link) -> Optional[Dict[int, str]]:
    """解析单个外部链接"""
    try:
        link_id = getattr(link, 'id', None)
        target = getattr(link, 'target', None) or getattr(link, 'file_link', None)

        if link_id is not None and target:
            if isinstance(link_id, int):
                return {link_id: target}

        # 尝试从字符串表示中提取信息
        link_str = str(link)
        id_match = re.search(r'id=(\d+)', link_str)
        target_match = re.search(r"target='([^']+)'", link_str)

        if id_match and target_match:
            return {int(id_match.group(1)): target_match.group(1)}

    except Exception:
        pass

    return None


# ==================== 模板结构读取函数 ====================

def read_template_structure(
    template_file: str,
    template_sheet: str
) -> Tuple[List[str], Dict[str, str], openpyxl.worksheet.worksheet.Worksheet, openpyxl.Workbook]:
    """
    读取模板的结构信息

    Args:
        template_file: 模板文件路径
        template_sheet: 模板sheet名称

    Returns:
        Tuple: (列名列表, 公式模板字典, 模板工作表对象, 工作簿对象)
               注意：调用者需要负责关闭返回的工作簿对象
    """
    wb = openpyxl.load_workbook(template_file, data_only=False)

    try:
        if template_sheet not in wb.sheetnames:
            raise ValueError(f"模板中不存在工作表: {template_sheet}")

        ws = wb[template_sheet]

        # 读取第一行作为列名
        columns = _read_template_columns(ws)

        # 读取第二行的公式（作为公式模板）
        formula_templates = _read_formula_templates(ws, columns)

        print(f"正在读取模板: {os.path.basename(template_file)} 的 {template_sheet} 工作表...")
        print(f"模板列名: {columns}")
        print(f"公式列: {[k for k, v in formula_templates.items() if v]}")

        return columns, formula_templates, ws, wb
    except Exception:
        wb.close()
        raise


def _read_template_columns(ws) -> List[str]:
    """读取模板列名，同名列自动添加后缀 _2, _3 ..."""
    seen: Dict[str, int] = {}
    columns = []
    for cell in ws[1]:
        if cell.value:
            name = str(cell.value)
            count = seen.get(name, 0)
            seen[name] = count + 1
            if count > 0:
                name = f"{name}_{count + 1}"
            columns.append(name)
        else:
            columns.append(f"Col_{len(columns) + 1}")
    return columns


def _read_formula_templates(ws, columns: List[str]) -> Dict[str, str]:
    """读取公式模板"""
    formula_templates = {}

    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=2, column=col_idx)
        if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
            formula_templates[col_name] = cell.value

    return formula_templates


# ==================== 公式替换辅助函数 ====================

def replace_link_indices_with_filenames(formula: str, link_mapping: Dict[int, str]) -> str:
    """
    将公式中的链接索引号替换为实际文件名

    Args:
        formula: 原始公式字符串
        link_mapping: 链接映射 {索引号: 文件名}

    Returns:
        str: 替换后的公式
    """
    if not link_mapping:
        return formula

    # 匹配 [数字] 格式的链接索引
    pattern = r'\[(\d+)\]'

    def replace_index(match):
        index = int(match.group(1))
        if index in link_mapping:
            filename = os.path.basename(link_mapping[index])
            return f'[{filename}]'
        return match.group(0)

    return re.sub(pattern, replace_index, formula)


# ==================== 公式引用解析函数 ====================

def parse_formula_references(formula: str) -> List[Tuple[str, str, int]]:
    """
    解析公式中的引用

    Args:
        formula: 公式字符串

    Returns:
        List[Tuple]: [(sheet名, 列引用, 行号), ...]
    """
    references = []

    # 匹配单元格引用的正则模式
    patterns = [
        r"'([^']+)'!([A-Z]+)(\d+)",  # 'SheetName'!A1
        r"([A-Za-z_][A-Za-z0-9_]*)!([A-Z]+)(\d+)",  # SheetName!A1
    ]

    for pattern in patterns:
        matches = re.findall(pattern, formula)
        for match in matches:
            # 将行号转换为整数
            sheet, col, row = match
            references.append((sheet, col, int(row)))

    return references


# ==================== Sheet引用替换函数 ====================

# 正则模式常量
# 注意: 添加 \$? 来支持绝对引用（如 $A$1, $D:$D 等）
# 注意: CJK 范围 一-鿿 支持中文等非 ASCII 表名(Excel 不给中文加引号)
_QUOTED_PATTERN = r"'([^']+)'!(\$?[A-Z]+\$?\d*(?::\$?[A-Z]*\$?\d*)?)"
_BRACKET_PATTERN = r"(\[[^\]]+\][^!'\s]+)!(\$?[A-Z]+\$?\d*(?::\$?[A-Z]*\$?\d*)?)"
_UNQUOTED_PATTERN = r"([A-Za-z_一-鿿][A-Za-z0-9_.一-鿿]*)!(\$?[A-Z]+\$?\d*(?::\$?[A-Z]*\$?\d*)?)"
_LOCAL_PATTERN = r"(?<![A-Za-z!'\"\\])(\$?[A-Z]+)(\$?\d+)(?![A-Za-z])"


def replace_sheet_references(
    formula: str,
    alias_to_info: Dict[str, Dict[str, str]],
    row_offset: int = 0,
    output_file_path: Optional[str] = None,
    external_links: Optional[Dict[int, str]] = None
) -> str:
    """
    替换公式中的sheet引用为外部文件引用，并调整行号
    """
    # 先处理带单引号的sheet名
    result = re.sub(_QUOTED_PATTERN,
                    lambda m: _replace_quoted_match(m, alias_to_info, row_offset, output_file_path, external_links),
                    formula, flags=re.IGNORECASE)

    # 处理带方括号但无单引号的格式
    result = re.sub(_BRACKET_PATTERN,
                    lambda m: _replace_bracket_match(m, alias_to_info, row_offset, output_file_path, external_links),
                    result, flags=re.IGNORECASE)

    # 最后处理不带单引号和方括号的sheet名
    result = re.sub(_UNQUOTED_PATTERN,
                    lambda m: _replace_unquoted_match(m, alias_to_info, row_offset, output_file_path),
                    result, flags=re.IGNORECASE)

    # 调整本地单元格引用的行号
    def _adjust_local_ref(m, offset):
        col_part = m.group(1)
        row_part = m.group(2)
        if row_part.startswith('$'):
            return f"{col_part}{row_part}"
        return f"{col_part}{int(row_part) + offset}"

    result = re.sub(_LOCAL_PATTERN,
                    lambda m: _adjust_local_ref(m, row_offset),
                    result)

    return result


def _adjust_cell_ref(cell_ref: str, row_offset: int) -> str:
    """调整单元格引用的行号"""
    if ':' in cell_ref:
        return _adjust_range_ref(cell_ref, row_offset)
    return _adjust_single_ref(cell_ref, row_offset)


def _adjust_single_ref(cell_ref: str, row_offset: int) -> str:
    m = re.match(r'(\$?)([A-Z]+)(\$?)(\d+)', cell_ref)
    if m:
        col_dollar, col, row_dollar, row_str = m.group(1), m.group(2), m.group(3), m.group(4)
        row = int(row_str) if row_dollar == '$' else int(row_str) + row_offset
        return f"{col_dollar}{col}{row_dollar}{row}"
    return cell_ref


def _adjust_range_ref(cell_ref: str, row_offset: int) -> str:
    parts = cell_ref.split(':')
    adjusted_parts = []
    for part in parts:
        m = re.match(r'(\$?)([A-Z]+)(\$?)(\d*)', part)
        if m:
            col_dollar, col, row_dollar, row_str = m.group(1), m.group(2), m.group(3), m.group(4)
            if row_str:
                row = int(row_str) if row_dollar == '$' else int(row_str) + row_offset
                adjusted_parts.append(f"{col_dollar}{col}{row_dollar}{row}")
            else:
                adjusted_parts.append(f"{col_dollar}{col}")
        else:
            adjusted_parts.append(part)
    return ':'.join(adjusted_parts)


def _extract_sheet_name(full_reference: str) -> str:
    """从完整引用中提取实际的sheet名"""
    bracket_match = re.match(r'\[.+\](.+)', full_reference)
    if bracket_match:
        return bracket_match.group(1)
    return full_reference


def _find_matching_info(sheet_name: str, alias_to_info: Dict, external_links: Optional[Dict]) -> Optional[Dict]:
    """查找匹配的sheet信息"""
    actual_sheet_name = _extract_sheet_name(sheet_name)
    actual_sheet_name_lower = actual_sheet_name.lower()

    # 1. 精确匹配
    for key, info in alias_to_info.items():
        if key.lower() == actual_sheet_name_lower:
            return info

    # 2. 匹配实际的sheet名
    matching_infos = [
        (key, info) for key, info in alias_to_info.items()
        if info.get('sheet_name', '').lower() == actual_sheet_name_lower
    ]

    if len(matching_infos) > 1:
        return _resolve_multiple_matches(sheet_name, matching_infos)
    elif matching_infos:
        return matching_infos[0][1]

    # 3. 处理数字索引
    return _find_by_index(sheet_name, alias_to_info, external_links)


def _resolve_multiple_matches(sheet_name: str, matching_infos: List) -> Optional[Dict]:
    """解决多个匹配的情况"""
    filename_match = re.search(r'\[([^\]]+)\]', sheet_name)
    if filename_match:
        filename = filename_match.group(1).lower()
        for key, info in matching_infos:
            if filename in os.path.basename(info['file_path']).lower():
                return info
    return matching_infos[0][1]


def _find_by_index(sheet_name: str, alias_to_info: Dict, external_links: Optional[Dict]) -> Optional[Dict]:
    """通过索引查找"""
    bracket_match = re.match(r'\[(\d+)\]', sheet_name)
    if not bracket_match:
        return None

    index = bracket_match.group(1)

    if external_links and index in external_links:
        external_filename = external_links[index].lower()
        for key, info in alias_to_info.items():
            if external_filename in os.path.basename(info['file_path']).lower():
                return info

    return alias_to_info.get(index)


def _build_reference(info: Dict, adjusted_ref: str, output_file_path: Optional[str]) -> str:
    """构建引用字符串"""
    file_path = info.get('file_path', '')
    actual_sheet_name = info['sheet_name']

    # 检查是否为内部引用（数据源sheet已复制到输出文件）
    is_internal = info.get('is_internal', False)

    # 检查是否为本地引用
    is_local = (
        is_internal or
        (output_file_path and file_path and os.path.normpath(file_path) == os.path.normpath(output_file_path)) or
        info.get('is_template_self_reference', False)
    )

    if is_local:
        if any(c in actual_sheet_name for c in " -()&^%$#@!~`'\"\\."):
            return f"'{actual_sheet_name}'!{adjusted_ref}"
        return f"{actual_sheet_name}!{adjusted_ref}"

    file_name = os.path.basename(file_path)
    return f"'[{file_name}]{actual_sheet_name}'!{adjusted_ref}"


def _replace_quoted_match(match, alias_to_info: Dict, row_offset: int,
                          output_file_path: Optional[str], external_links: Optional[Dict]) -> str:
    """替换带引号的匹配"""
    full_reference = match.group(1)
    cell_ref = match.group(2).upper()

    info = _find_matching_info(full_reference, alias_to_info, external_links)
    if info:
        adjusted_ref = _adjust_cell_ref(cell_ref, row_offset)
        return _build_reference(info, adjusted_ref, output_file_path)
    return match.group(0)


def _replace_unquoted_match(match, alias_to_info: Dict, row_offset: int,
                            output_file_path: Optional[str]) -> str:
    """替换不带引号的匹配"""
    sheet_name = match.group(1)
    cell_ref = match.group(2).upper()

    info = _find_matching_info(sheet_name, alias_to_info, None)
    if info:
        adjusted_ref = _adjust_cell_ref(cell_ref, row_offset)
        return _build_reference(info, adjusted_ref, output_file_path)
    return match.group(0)


def _replace_bracket_match(match, alias_to_info: Dict, row_offset: int,
                           output_file_path: Optional[str], external_links: Optional[Dict]) -> str:
    """替换带方括号的匹配"""
    full_reference = match.group(1)
    cell_ref = match.group(2).upper()

    info = _find_matching_info(full_reference, alias_to_info, external_links)
    if info:
        adjusted_ref = _adjust_cell_ref(cell_ref, row_offset)
        return _build_reference(info, adjusted_ref, output_file_path)
    return match.group(0)


# ==================== 工作表复制函数 ====================

def _copy_worksheet(source_ws, target_ws) -> None:
    """
    复制工作表的所有内容（数据和样式）

    Args:
        source_ws: 源工作表
        target_ws: 目标工作表
    """
    # 复制单元格数据和样式
    for row in source_ws.iter_rows():
        for cell in row:
            new_cell = target_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                copy_cell_style(cell, new_cell)

    # 复制合并单元格
    for merged_range in source_ws.merged_cells.ranges:
        target_ws.merge_cells(str(merged_range))

    # 复制列宽
    for col_letter, col_dim in source_ws.column_dimensions.items():
        if col_dim.width:
            target_ws.column_dimensions[col_letter].width = col_dim.width

    # 复制行高
    for row_idx, row_dim in source_ws.row_dimensions.items():
        if row_dim.height:
            target_ws.row_dimensions[row_idx].height = row_dim.height
